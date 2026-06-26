"""Generate detailed manual/integration QA Excel for pending backend EC-AUTO tests."""
from __future__ import annotations

import importlib.util
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUTPUT = (
    Path(__file__).resolve().parent.parent
    / "docs"
    / "EC-AUTO-Backend-Pending.xlsx"
)

BASE_URL = "http://127.0.0.1:3002"
LOCAL_DB = "postgresql://fineset:fineset@localhost:5432/fineset_test"

HEADERS = [
    "Test Case ID",
    "API / Component (Where to Test)",
    "HTTP Method & Endpoint",
    "Problem Description (What Is Wrong or Unverified)",
    "Category",
    "Feature Area",
    "Priority",
    "Severity",
    "Action Required",
    "Matrix Status",
    "Auth / Role Required",
    "Environment / Setup",
    "How to Test (Step by Step)",
    "Expected Outcome",
    "Pass Criteria (What Should Happen)",
    "Watch For During Test",
    "Actual Result (Fill In)",
    "Pass / Fail (Fill In)",
    "Tester Name (Fill In)",
    "Test Date (Fill In)",
    "Notes (Fill In)",
]

PRIORITY_FILLS = {
    "P1": PatternFill("solid", fgColor="FFC7CE"),
    "P2": PatternFill("solid", fgColor="FFEB9C"),
    "P3": PatternFill("solid", fgColor="E2EFDA"),
}

# api: human label | endpoint: method path | problem: gap or verify note
CASE_BACKEND: dict[str, dict[str, str]] = {
    "EC-AUTO-003": {
        "api": "Admin automation config API",
        "endpoint": "GET /api/admin/automation/config",
        "problem": "Unauthenticated requests must return 401. Vitest: automation-backend.test.ts EC-AUTO-003.",
    },
    "EC-AUTO-006": {
        "api": "Admin automation config API",
        "endpoint": "PATCH /api/admin/automation/config",
        "problem": "Platform admin must not PATCH config (403). Vitest: EC-AUTO-006.",
    },
    "EC-AUTO-007": {
        "api": "Admin automation run API",
        "endpoint": "POST /api/admin/automation/run",
        "problem": "Platform admin must not trigger runs (403). Vitest: EC-AUTO-007.",
    },
    "EC-AUTO-010": {
        "api": "Cron billing automation",
        "endpoint": "GET /api/cron/billing-automation",
        "problem": "Cron must return 503 when CRON_SECRET env unset. Vitest: EC-AUTO-010.",
    },
    "EC-AUTO-011": {
        "api": "Cron billing automation",
        "endpoint": "GET /api/cron/billing-automation",
        "problem": "Wrong Bearer token must return 401. Vitest: EC-AUTO-117 (same endpoint).",
    },
    "EC-AUTO-012": {
        "api": "Cron billing automation",
        "endpoint": "GET /api/cron/billing-automation",
        "problem": "Valid Bearer must trigger CRON run and create run log. Not integration-tested yet.",
    },
    "EC-AUTO-045": {
        "api": "Admin automation config API",
        "endpoint": "PATCH /api/admin/automation/config",
        "problem": "Empty PATCH body {} should no-op and return current config unchanged.",
    },
    "EC-AUTO-047": {
        "api": "Admin automation config API",
        "endpoint": "PATCH /api/admin/automation/config",
        "problem": "Unknown extra JSON keys must be stripped/ignored by Zod schema.",
    },
    "EC-AUTO-048": {
        "api": "Admin automation run API",
        "endpoint": "POST /api/admin/automation/run",
        "problem": "Invalid dryRun type in body must return 400 validation error.",
    },
    "EC-AUTO-051": {
        "api": "Admin automation runs list API",
        "endpoint": "GET /api/admin/automation/runs?page=0&pageSize=100",
        "problem": "Invalid page/pageSize must clamp to safe range (not 500 error).",
    },
    "EC-AUTO-052": {
        "api": "Admin automation runs list API",
        "endpoint": "GET /api/admin/automation/runs?page=abc",
        "problem": "Non-numeric page param must default to page 1.",
    },
    "EC-AUTO-058": {
        "api": "Automation config service",
        "endpoint": "GET /api/admin/automation/config (fresh DB)",
        "problem": "First load with no DB row must return merged default config.",
    },
    "EC-AUTO-059": {
        "api": "Automation config service",
        "endpoint": "GET /api/admin/automation/config",
        "problem": "Corrupt partial JSON in DB must merge safely with defaults.",
    },
    "EC-AUTO-060": {
        "api": "Automation config service",
        "endpoint": "GET /api/admin/automation/config",
        "problem": "Non-object config JSONB in DB must fall back to full defaults.",
    },
    "EC-AUTO-061": {
        "api": "Automation config cache",
        "endpoint": "GET /api/admin/automation/config (rapid reads)",
        "problem": "30s cache TTL — verify fresh read after PATCH within cache window.",
    },
    "EC-AUTO-063": {
        "api": "Admin automation runs list API",
        "endpoint": "GET /api/admin/automation/runs",
        "problem": "Malformed summary/errors JSON on run log must not crash list API.",
    },
    "EC-AUTO-065": {
        "api": "Billing automation runner — dedupe",
        "endpoint": "POST /api/admin/automation/run (twice same day)",
        "problem": "Same business+action same cycle must skip second send via dedupeKey.",
    },
    "EC-AUTO-066": {
        "api": "Billing automation runner — dry run",
        "endpoint": "POST /api/admin/automation/run {\"dryRun\":true}",
        "problem": "Dry run must not write AutomationDeliveryLog rows.",
    },
    "EC-AUTO-067": {
        "api": "Billing automation runner — dry run",
        "endpoint": "POST /api/admin/automation/run {\"dryRun\":true}",
        "problem": "Dry run must bypass dedupe check (by design) — verify no false skips.",
    },
    "EC-AUTO-068": {
        "api": "run-billing-automation.ts",
        "endpoint": "POST /api/admin/automation/run",
        "problem": "If config load throws, runner must fall back to default billing settings.",
    },
    "EC-AUTO-069": {
        "api": "Cron + runner (global disabled)",
        "endpoint": "GET /api/cron/billing-automation",
        "problem": "When global.enabled=false, cron run log must be SUCCESS with skipped detail.",
    },
    "EC-AUTO-072": {
        "api": "Runner — invoice phase",
        "endpoint": "POST /api/admin/automation/run",
        "problem": "SMTP missing must skip invoice phase with error logged, not crash run.",
    },
    "EC-AUTO-073": {
        "api": "Runner — invoice phase",
        "endpoint": "POST /api/admin/automation/run",
        "problem": "Business without email must skip invoice with detail message.",
    },
    "EC-AUTO-074": {
        "api": "Runner — invoice phase",
        "endpoint": "POST /api/admin/automation/run",
        "problem": "Paid/CURRENT account with skipIfPaid=true must not receive invoice.",
    },
    "EC-AUTO-075": {
        "api": "Runner — invoice dedupe",
        "endpoint": "POST /api/admin/automation/run (same cycle twice)",
        "problem": "Second invoice attempt same billing cycle must dedupe skip.",
    },
    "EC-AUTO-076": {
        "api": "Runner — invoice schedule",
        "endpoint": "POST /api/admin/automation/run",
        "problem": "Invoice must send when today matches sendDayOfMonth / renewal schedule.",
    },
    "EC-AUTO-077": {
        "api": "Runner — invoice phase",
        "endpoint": "POST /api/admin/automation/run",
        "problem": "Non-invoice day must skip businesses silently (no send attempts).",
    },
    "EC-AUTO-078": {
        "api": "Runner — payment reminders",
        "endpoint": "POST /api/admin/automation/run",
        "problem": "paymentReminders.enabled=false must skip entire reminder phase.",
    },
    "EC-AUTO-079": {
        "api": "Runner — payment reminders",
        "endpoint": "POST /api/admin/automation/run",
        "problem": "Today not in before/after due day lists — no reminders sent.",
    },
    "EC-AUTO-081": {
        "api": "Runner — payment reminders",
        "endpoint": "POST /api/admin/automation/run",
        "problem": "PAID/WAIVED accounts with stopAfterPayment=true must not get reminders.",
    },
    "EC-AUTO-082": {
        "api": "Runner — payment reminders",
        "endpoint": "POST /api/admin/automation/run",
        "problem": "Max reminders per cycle reached — skip with detail in run summary.",
    },
    "EC-AUTO-083": {
        "api": "Runner — reminders email + WhatsApp",
        "endpoint": "POST /api/admin/automation/run",
        "problem": "Email failure must not block WhatsApp attempt (independent channels).",
    },
    "EC-AUTO-084": {
        "api": "Runner — WhatsApp reminders",
        "endpoint": "POST /api/admin/automation/run",
        "problem": "WhatsApp only when BOTH paymentReminders.whatsAppEnabled AND whatsApp.enabled.",
    },
    "EC-AUTO-085": {
        "api": "Runner — WhatsApp business hours",
        "endpoint": "POST /api/admin/automation/run",
        "problem": "businessHoursOnly=true + run outside hours — WhatsApp skipped.",
    },
    "EC-AUTO-087": {
        "api": "Runner — follow-ups",
        "endpoint": "POST /api/admin/automation/run",
        "problem": "Eligible unpaid accounts must get follow-ups auto-scheduled when enabled.",
    },
    "EC-AUTO-088": {
        "api": "Runner — follow-ups escalation",
        "endpoint": "POST /api/admin/automation/run",
        "problem": "Max follow-ups + escalateAfterMax=true must log for manual review.",
    },
    "EC-AUTO-089": {
        "api": "Runner — follow-up spacing",
        "endpoint": "POST /api/admin/automation/run",
        "problem": "spacingDays shorter than count must reuse last spacing value.",
    },
    "EC-AUTO-090": {
        "api": "Runner — renewal/expiry reminders",
        "endpoint": "POST /api/admin/automation/run",
        "problem": "Renewal/expiry emails must fire on matching offset days only.",
    },
    "EC-AUTO-091": {
        "api": "Runner — monthly reports",
        "endpoint": "POST /api/admin/automation/run",
        "problem": "Wrong day/hour for monthly report — phase skipped (no send).",
    },
    "EC-AUTO-092": {
        "api": "Runner — monthly reports",
        "endpoint": "POST /api/admin/automation/run",
        "problem": ">50 businesses + includePerStoreMetrics — report truncates with 'and N more'.",
    },
    "EC-AUTO-095": {
        "api": "Runner — run status logic",
        "endpoint": "POST /api/admin/automation/run",
        "problem": "Mixed success/failure must set status SUCCESS / PARTIAL / FAILED correctly.",
    },
    "EC-AUTO-096": {
        "api": "Runner — error handling",
        "endpoint": "POST /api/admin/automation/run",
        "problem": "Uncaught fatal error during run must mark AutomationRunLog as FAILED.",
    },
    "EC-AUTO-097": {
        "api": "Payment confirmation email",
        "endpoint": "Mark billing account PAID (billing API)",
        "problem": "paymentConfirmationEnabled=true + SMTP — confirmation email sent on PAID.",
    },
    "EC-AUTO-098": {
        "api": "Payment confirmation email",
        "endpoint": "Mark billing account PAID (billing API)",
        "problem": "paymentConfirmationEnabled=false — no confirmation email on PAID.",
    },
    "EC-AUTO-099": {
        "api": "Payment confirmation dedupe",
        "endpoint": "Mark billing account PAID twice",
        "problem": "Second PAID in same cycle must not resend confirmation (dedupe).",
    },
    "EC-AUTO-100": {
        "api": "Billing payment side effects",
        "endpoint": "Record payment on billing account",
        "problem": "autoExtendOnPayment=false — renewal dates must not change on payment.",
    },
    "EC-AUTO-102": {
        "api": "Runner — all email phases",
        "endpoint": "POST /api/admin/automation/run",
        "problem": "Missing SMTP env — all email phases skip gracefully without crash.",
    },
    "EC-AUTO-103": {
        "api": "automation-emails.ts templates",
        "endpoint": "POST /api/admin/automation/run (reminder variants)",
        "problem": "Email copy must differ for overdue vs due-today vs upcoming reminders.",
    },
    "EC-AUTO-104": {
        "api": "Payment confirmation template",
        "endpoint": "Trigger payment confirmation",
        "problem": "No invoice number — confirmation email omits invoice line.",
    },
    "EC-AUTO-107": {
        "api": "Invoice email sender",
        "endpoint": "POST /api/admin/automation/run (invoice day)",
        "problem": "Invoice From address must match platform automation sender.",
    },
    "EC-AUTO-108": {
        "api": "Email templates — branding",
        "endpoint": "POST /api/admin/automation/run",
        "problem": "Unset support email — support line omitted from template.",
    },
    "EC-AUTO-109": {
        "api": "Runner — partial failure",
        "endpoint": "POST /api/admin/automation/run",
        "problem": "Individual send failure (bad recipient) — run status PARTIAL not FAILED.",
    },
    "EC-AUTO-112": {
        "api": "Automation config snapshot",
        "endpoint": "POST /api/admin/automation/run + PATCH config mid-run",
        "problem": "Active run must use config loaded at start, not mid-run PATCH.",
    },
    "EC-AUTO-115": {
        "api": "Automation delivery dedupe keys",
        "endpoint": "Invoice + payment confirmation same cycle",
        "problem": "Separate dedupe keys — both can send without blocking each other.",
    },
    "EC-AUTO-116": {
        "api": "AutomationDeliveryLog upsert",
        "endpoint": "Concurrent POST /api/admin/automation/run",
        "problem": "Parallel delivery log writes same dedupeKey — no duplicate rows.",
    },
    "EC-AUTO-119": {
        "api": "Runner — monthly report recipients",
        "endpoint": "POST /api/admin/automation/run",
        "problem": "recipients=business_owners only sends when explicitly configured — scope review.",
    },
    "EC-AUTO-120": {
        "api": "Admin automation APIs — CSRF",
        "endpoint": "PATCH/POST from cross-origin",
        "problem": "Forged cross-origin PATCH/POST blocked by session cookie policy.",
    },
    "EC-AUTO-122": {
        "api": "Runner — empty portfolio",
        "endpoint": "POST /api/admin/automation/run",
        "problem": "Zero businesses — run completes SUCCESS with all counts zero.",
    },
    "EC-AUTO-123": {
        "api": "Runner — all phases fail",
        "endpoint": "POST /api/admin/automation/run",
        "problem": "All actions error — final run status must be FAILED.",
    },
    "EC-AUTO-125": {
        "api": "Cron + runner (disabled)",
        "endpoint": "GET /api/cron/billing-automation",
        "problem": "global.enabled=false on cron — SUCCESS log with 'automations disabled' detail.",
    },
}


def load_matrix_rows():
    matrix_path = Path(__file__).resolve().parent / "generate-automation-qa-excel.py"
    spec = importlib.util.spec_from_file_location("automation_matrix", matrix_path)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module.ROWS


def layer_from_module(module: str) -> str:
    if "(Backend)" in module:
        return "Backend"
    return "Other"


def feature_area(module: str) -> str:
    return module.split("(")[0].strip()


def is_backend_pending(row: dict) -> bool:
    return layer_from_module(row["module"]) == "Backend" and row["status"] == "Not Tested"


def auth_required(row: dict) -> str:
    text = f"{row['pre']} {row['steps']}".upper()
    if "CRON" in text or "BEARER" in text:
        return "Bearer CRON_SECRET (cron) or none for 503 test"
    if "PLATFORM_ADMIN" in text:
        return "Platform Admin session cookie"
    if "NO SESSION" in text or "WITHOUT AUTH" in text:
        return "No session (unauthenticated)"
    return "Master Admin session cookie"


def backend_details(row: dict) -> tuple[str, str, str]:
    case_id = row["id"]
    if case_id in CASE_BACKEND:
        entry = CASE_BACKEND[case_id]
        return entry["api"], entry["endpoint"], entry["problem"]
    return (
        feature_area(row["module"]),
        "See test steps",
        row["remarks"] or f"Not verified: {row['scenario']}",
    )


def how_to_test(row: dict, endpoint: str) -> str:
    case_id = row["id"]
    lines = [
        "⚠ Use LOCAL Postgres only — do NOT point at Supabase staging.",
        f"DATABASE_URL={LOCAL_DB}",
        "",
    ]

    vitest_file = "tests/integration/automation-backend-pending.test.ts"
    legacy_file = "tests/integration/automation-backend.test.ts"

    if "Vitest" in row["tool"] or case_id in {
        "EC-AUTO-003", "EC-AUTO-006", "EC-AUTO-007", "EC-AUTO-010", "EC-AUTO-011",
    }:
        lines.extend([
            "Option A — Automated (recommended, LOCAL Postgres only):",
            f"  npx vitest run {vitest_file} -t \"{case_id}\"",
            f"  # or legacy: npx vitest run {legacy_file} -t \"{case_id}\"",
            "",
            "Option B — Manual API (curl / Postman):",
        ])
    elif row["tool"] == "Vitest unit":
        lines.extend([
            "Run unit test:",
            f"  npx vitest run -t \"{row['scenario'][:40]}\"",
            "",
        ])
    else:
        lines.append(
            f"Integration test via Vitest ({vitest_file}) + local Postgres (recommended)."
        )

    if endpoint.startswith("GET /api/cron"):
        lines.extend([
            f"  curl -i \"{BASE_URL}{endpoint.split(' ', 1)[1].split('?')[0]}\" \\",
            "    -H \"Authorization: Bearer $CRON_SECRET\"",
        ])
    elif "PATCH" in endpoint:
        path = endpoint.split(" ", 1)[1].split("?")[0]
        lines.extend([
            f"  curl -i -X PATCH \"{BASE_URL}{path}\" \\",
            "    -H \"Content-Type: application/json\" \\",
            "    -H \"Cookie: <master-admin-session>\" \\",
            "    -d '{\"global\":{\"enabled\":true}}'",
        ])
    elif "POST /api/admin/automation/run" in endpoint:
        lines.extend([
            f"  curl -i -X POST \"{BASE_URL}/api/admin/automation/run\" \\",
            "    -H \"Content-Type: application/json\" \\",
            "    -H \"Cookie: <master-admin-session>\" \\",
            "    -d '{\"dryRun\":true}'",
        ])
    elif endpoint.startswith("GET /api"):
        path = endpoint.split(" ", 1)[1]
        lines.append(f"  curl -i \"{BASE_URL}{path}\" -H \"Cookie: <session>\"")
    else:
        for step in row["steps"].split("\n"):
            lines.append(f"  {step.strip()}")

    lines.extend(["", "Matrix steps:", row["steps"]])
    return "\n".join(lines)


def pass_criteria(row: dict) -> str:
    return (
        f"{row['expected']}\n\n"
        f"Verify in: DB tables AutomationRunLog, AutomationDeliveryLog, "
        f"and API response body/status code."
    )


def environment_setup(row: dict) -> str:
    parts = [
        f"Local dev server: {BASE_URL}",
        f"Local test DB: {LOCAL_DB}",
        row["pre"],
    ]
    pre = row["pre"].lower()
    if "cron" in pre:
        parts.append("Set CRON_SECRET in .env for positive cron tests; unset for EC-AUTO-010.")
    if "smtp" in pre:
        parts.append("Configure or intentionally omit SMTP_* env vars per test.")
    if "seed" in pre or "business" in pre:
        parts.append("Seed billing businesses via fixtures or perf-seed on LOCAL DB only.")
    return "\n".join(p for p in parts if p)


def row_to_values(row: dict) -> list:
    api, endpoint, problem = backend_details(row)
    return [
        row["id"],
        api,
        endpoint,
        problem,
        row["module"],
        feature_area(row["module"]),
        row["priority"],
        row["severity"],
        "Execute now",
        row["status"],
        auth_required(row),
        environment_setup(row),
        how_to_test(row, endpoint),
        row["expected"],
        pass_criteria(row),
        row["auto_notes"] or "Record HTTP status, response JSON, and DB rows.",
        "",
        "",
        "",
        "",
        "",
    ]


def write_data_sheet(ws, rows: list[dict], title_note: str = "") -> None:
    start_row = 3 if title_note else 1
    if title_note:
        ws.cell(row=1, column=1, value=title_note).font = Font(bold=True, size=12)

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    for col, header in enumerate(HEADERS, 1):
        cell = ws.cell(row=start_row, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for idx, row in enumerate(rows, start_row + 1):
        for col_idx, value in enumerate(row_to_values(row), 1):
            cell = ws.cell(row=idx, column=col_idx, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if col_idx == 7:
                cell.fill = PRIORITY_FILLS.get(row["priority"], PatternFill())

    ws.freeze_panes = ws.cell(row=start_row + 1, column=1).coordinate
    if rows:
        ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(HEADERS))}{start_row + len(rows)}"

    widths = [14, 34, 38, 44, 32, 24, 8, 10, 12, 12, 28, 38, 48, 36, 40, 32, 20, 12, 16, 14, 24]
    for col_idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def write_instructions(ws) -> None:
    lines = [
        ("Automation Center — Backend Pending QA (60 cases)", 14, True),
        ("", 11, False),
        ("Purpose", 12, True),
        (
            "Execute all 60 backend tests marked Not Tested. Each row shows API/component, "
            "problem to verify, steps, and expected outcome.",
            11,
            False,
        ),
        ("", 11, False),
        ("Safety", 12, True),
        (f"• Use LOCAL Postgres: {LOCAL_DB}", 11, False),
        ("• Do NOT run integration tests against Supabase staging (.env.local)", 11, False),
        ("• Prefer: npx vitest run tests/integration/automation-backend-pending.test.ts", 11, False),
        ("• Legacy cases also in: tests/integration/automation-backend.test.ts", 11, False),
        ("", 11, False),
        ("Key columns", 12, True),
        ("• API / Component — service or route under test", 11, False),
        ("• HTTP Method & Endpoint — exact API path", 11, False),
        ("• Problem Description — what is unverified or may be broken", 11, False),
        ("", 11, False),
        ("API map", 12, True),
        ("• GET/PATCH /api/admin/automation/config — settings", 11, False),
        ("• POST /api/admin/automation/run — manual/dry run", 11, False),
        ("• GET /api/admin/automation/runs — history list", 11, False),
        ("• GET /api/cron/billing-automation — scheduled cron", 11, False),
        ("• lib/services/run-billing-automation.ts — billing runner logic", 11, False),
    ]
    for row_idx, (text, size, bold) in enumerate(lines, 1):
        cell = ws.cell(row=row_idx, column=1, value=text)
        cell.font = Font(bold=bold, size=size)
    ws.column_dimensions["A"].width = 95


def main() -> None:
    all_rows = load_matrix_rows()
    pending = [r for r in all_rows if is_backend_pending(r)]
    pending.sort(key=lambda r: r["id"])

    wb = Workbook()
    inst = wb.active
    inst.title = "How To Use"
    write_instructions(inst)

    summary = wb.create_sheet("Summary")
    summary.cell(row=1, column=1, value="Backend Not Tested summary").font = Font(bold=True, size=14)
    summary.cell(row=3, column=1, value="Total pending backend cases").font = Font(bold=True)
    summary.cell(row=3, column=2, value=len(pending))

    by_area: dict[str, int] = {}
    for row in pending:
        area = feature_area(row["module"])
        by_area[area] = by_area.get(area, 0) + 1
    r = 5
    summary.cell(row=r, column=1, value="By feature area").font = Font(bold=True)
    r += 1
    for area, count in sorted(by_area.items()):
        summary.cell(row=r, column=1, value=area)
        summary.cell(row=r, column=2, value=count)
        r += 1

    all_sheet = wb.create_sheet("All 60 Pending")
    write_data_sheet(all_sheet, pending, f"ALL BACKEND NOT TESTED — {len(pending)} cases")

    for area in sorted(by_area):
        area_rows = [row for row in pending if feature_area(row["module"]) == area]
        safe = area[:28].replace("/", "-")
        write_data_sheet(wb.create_sheet(safe), area_rows)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)
    print(f"Wrote {OUTPUT} ({len(pending)} backend pending cases)")


if __name__ == "__main__":
    main()
