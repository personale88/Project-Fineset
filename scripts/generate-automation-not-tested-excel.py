"""Export EC-AUTO cases with Status = Not Tested to Excel."""
from __future__ import annotations

import importlib.util
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUTPUT = (
    Path(__file__).resolve().parent.parent / "docs" / "EC-AUTO-Not-Tested.xlsx"
)

HEADERS = [
    "Edge Case ID",
    "Module/Feature",
    "Layer",
    "Scenario Description",
    "Preconditions",
    "Test Steps",
    "Expected Result",
    "Code Review",
    "Actual Result",
    "Priority",
    "Severity",
    "Status",
    "Test Type",
    "Automation Tool",
    "Automation Notes",
    "Remarks",
]

TYPE_FILLS = {
    "Auto": PatternFill("solid", fgColor="C6EFCE"),
    "Manual": PatternFill("solid", fgColor="FFEB9C"),
    "Both": PatternFill("solid", fgColor="BDD7EE"),
}

LAYER_FILLS = {
    "Backend": PatternFill("solid", fgColor="E2EFDA"),
    "Frontend": PatternFill("solid", fgColor="FCE4D6"),
    "Full-stack": PatternFill("solid", fgColor="DDEBF7"),
}


def load_matrix_rows():
    matrix_path = Path(__file__).resolve().parent / "generate-automation-qa-excel.py"
    spec = importlib.util.spec_from_file_location("automation_matrix", matrix_path)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module.ROWS


def layer_from_module(module: str) -> str:
    if "(Backend)" in module:
        return "Backend"
    if "(Frontend)" in module:
        return "Frontend"
    if "(Full-stack)" in module:
        return "Full-stack"
    return "Other"


def write_sheet(ws, rows: list[dict]) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF", size=11)

    for col, header in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_idx, row in enumerate(rows, 2):
        layer = layer_from_module(row["module"])
        values = [
            row["id"],
            row["module"],
            layer,
            row["scenario"],
            row["pre"],
            row["steps"],
            row["expected"],
            row["review"],
            row["actual"],
            row["priority"],
            row["severity"],
            row["status"],
            row["type"],
            row["tool"],
            row["auto_notes"],
            row["remarks"],
        ]
        for col_idx, value in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if col_idx == 3:
                cell.fill = LAYER_FILLS.get(layer, PatternFill())
            if col_idx == 13:
                cell.fill = TYPE_FILLS.get(row["type"], PatternFill())

    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{len(rows) + 1}"

    widths = [14, 36, 12, 40, 32, 40, 32, 10, 36, 8, 10, 12, 10, 22, 28]
    for col_idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def write_summary(ws, not_tested: list[dict]) -> None:
    from collections import Counter

    by_layer = Counter(layer_from_module(r["module"]) for r in not_tested)
    by_type = Counter(r["type"] for r in not_tested)
    by_tool = Counter(r["tool"] for r in not_tested)

    ws.cell(row=1, column=1, value="EC-AUTO — Not Tested summary").font = Font(
        bold=True, size=14
    )
    ws.cell(row=3, column=1, value="Total Not Tested").font = Font(bold=True)
    ws.cell(row=3, column=2, value=len(not_tested))

    row = 5
    ws.cell(row=row, column=1, value="By layer").font = Font(bold=True)
    row += 1
    for label in ("Backend", "Frontend", "Full-stack", "Other"):
        if by_layer[label]:
            ws.cell(row=row, column=1, value=label)
            ws.cell(row=row, column=2, value=by_layer[label])
            row += 1

    row += 1
    ws.cell(row=row, column=1, value="By test type").font = Font(bold=True)
    row += 1
    for label, count in sorted(by_type.items(), key=lambda item: -item[1]):
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=2, value=count)
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="By automation tool").font = Font(bold=True)
    row += 1
    for label, count in sorted(by_tool.items(), key=lambda item: -item[1]):
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=2, value=count)
        row += 1

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 12


def main() -> None:
    all_rows = load_matrix_rows()
    not_tested = [row for row in all_rows if row["status"] == "Not Tested"]
    not_tested.sort(key=lambda row: row["id"])

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    write_summary(summary, not_tested)

    all_sheet = wb.create_sheet("All Not Tested")
    write_sheet(all_sheet, not_tested)

    for layer in ("Frontend", "Full-stack", "Backend"):
        layer_rows = [
            row for row in not_tested if layer_from_module(row["module"]) == layer
        ]
        if layer_rows:
            sheet = wb.create_sheet(layer)
            write_sheet(sheet, layer_rows)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)
    print(f"Wrote {OUTPUT} ({len(not_tested)} Not Tested cases)")


if __name__ == "__main__":
    main()
