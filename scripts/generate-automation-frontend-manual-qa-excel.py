"""Generate detailed manual QA Excel for Automation Center frontend tests."""
from __future__ import annotations

import importlib.util
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUTPUT = (
    Path(__file__).resolve().parent.parent
    / "docs"
    / "EC-AUTO-Frontend-Manual-QA.xlsx"
)

BASE_URL = "http://127.0.0.1:3002"
AUTOMATION_URL = f"{BASE_URL}/admin/dashboard/automation"

SCOPE_LABELS: dict[str, str] = {
    "overview": "Overview",
    "billingCycle": "Billing cycle",
    "invoices": "Invoices",
    "paymentReminders": "Payment reminders",
    "followUps": "Follow-ups",
    "expiryRenewal": "Expiry & renewal",
    "monthlyReports": "Monthly reports",
    "whatsApp": "WhatsApp",
    "history": "Run history",
}

# Screen + problem overrides for manual QA clarity.
CASE_SCREEN_AND_PROBLEM: dict[str, dict[str, str]] = {
    "EC-AUTO-001": {
        "screen": "Login page (before Automation Center)",
        "scope": "",
        "problem": "Unauthenticated users may reach admin routes without redirect to sign-in.",
    },
    "EC-AUTO-002": {
        "screen": "Login → Admin redirect",
        "scope": "",
        "problem": "Store owner/staff may access /admin/dashboard/automation without being blocked.",
    },
    "EC-AUTO-004": {
        "screen": "Automation Center — any tab (page-level)",
        "scope": "",
        "problem": "Was: page loaded but API returned 403 for platform admin without billing. Verify redirect to /admin/dashboard.",
    },
    "EC-AUTO-005": {
        "screen": "Automation Center — all config tabs",
        "scope": "overview",
        "problem": "Read-only platform admin may see editable controls or missing read-only hint.",
    },
    "EC-AUTO-008": {
        "screen": "Automation Center — any config section",
        "scope": "invoices",
        "problem": "Save may fail silently or not persist after refresh.",
    },
    "EC-AUTO-009": {
        "screen": "Overview + Run history",
        "scope": "overview",
        "problem": "Run Now may not show toast or new row in History after manual run.",
    },
    "EC-AUTO-013": {
        "screen": "Overview (Run Now button)",
        "scope": "overview",
        "problem": "Verify Run Now disabled when automations off; API must also block non-dry runs.",
    },
    "EC-AUTO-014": {
        "screen": "Login page / Automation Center",
        "scope": "",
        "problem": "Expired session may still show automation UI instead of login redirect.",
    },
    "EC-AUTO-015": {
        "screen": "Automation Center — initial page load",
        "scope": "overview",
        "problem": "No visible loading state on first visit — page may flash empty content.",
    },
    "EC-AUTO-016": {
        "screen": "Automation Center — page header / error banner",
        "scope": "overview",
        "problem": "When config API fails, user may not see error banner with Retry button.",
    },
    "EC-AUTO-017": {
        "screen": "Automation Center — all config tabs",
        "scope": "paymentReminders",
        "problem": "Platform admin may see enabled inputs instead of read-only disabled fields + hint.",
    },
    "EC-AUTO-018": {
        "screen": "Automation Center — left scope sidebar (mobile)",
        "scope": "overview",
        "problem": "On mobile, scope pills may not scroll horizontally.",
    },
    "EC-AUTO-019": {
        "screen": "Automation Center — scope sidebar + URL",
        "scope": "invoices",
        "problem": "Was: refresh always reset to Overview. Verify ?scope= keeps same tab.",
    },
    "EC-AUTO-020": {
        "screen": "Automation Center — any config tab",
        "scope": "invoices",
        "problem": "Was: unsaved draft lost when config refetches. Verify edits survive background refresh.",
    },
    "EC-AUTO-022": {
        "screen": "WhatsApp tab (business hours fields)",
        "scope": "whatsApp",
        "problem": "Invalid values show generic toast only — no inline field error under the input.",
    },
    "EC-AUTO-023": {
        "screen": "Overview (Run Now toast)",
        "scope": "overview",
        "problem": "Was: success toast missing renewals/expiry/reports counts. Verify full summary line.",
    },
    "EC-AUTO-024": {
        "screen": "Overview (Run Now toast)",
        "scope": "overview",
        "problem": "Was: failed run showed generic toast. Verify error details appear in toast description.",
    },
    "EC-AUTO-025": {
        "screen": "Overview (Run Now / Dry Run buttons)",
        "scope": "overview",
        "problem": "Run Now may stay enabled when global automations master switch is OFF.",
    },
    "EC-AUTO-027": {
        "screen": "Overview + all config tabs",
        "scope": "overview",
        "problem": "Read-only user may still see Save or Run buttons.",
    },
    "EC-AUTO-028": {
        "screen": "Run history tab",
        "scope": "history",
        "problem": "History tab may show blank area with no loading indicator while runs fetch.",
    },
    "EC-AUTO-029": {
        "screen": "Run history tab",
        "scope": "history",
        "problem": "Empty history may show no helpful empty-state message.",
    },
    "EC-AUTO-030": {
        "screen": "Run history tab",
        "scope": "history",
        "problem": "Was: only 20 runs shown with no Load more. Verify pagination button appears.",
    },
    "EC-AUTO-031": {
        "screen": "Run history tab",
        "scope": "history",
        "problem": "Was: History API failure showed no error UI. Verify banner + Retry.",
    },
    "EC-AUTO-032": {
        "screen": "Run history tab (status badge)",
        "scope": "history",
        "problem": "Was: RUNNING jobs looked like generic badge. Verify gold/distinct RUNNING style.",
    },
    "EC-AUTO-033": {
        "screen": "Run history tab (timestamps)",
        "scope": "history",
        "problem": "Was: timestamps hardcoded to India locale. Verify times match Overview timezone setting.",
    },
    "EC-AUTO-034": {
        "screen": "Payment reminders / Follow-ups / Expiry / Monthly reports tabs",
        "scope": "paymentReminders",
        "problem": "Was: 5 schema fields hidden from UI. Verify stopAfterPayment, escalateAfterMax, autoExtendOnPayment, includeBillingSummary, sendOnRenewalDue.",
    },
    "EC-AUTO-035": {
        "screen": "Invoices tab",
        "scope": "invoices",
        "problem": "Was: sendOnRenewalDue toggle missing. Verify toggle visible and saves.",
    },
    "EC-AUTO-036": {
        "screen": "WhatsApp tab + backend run (verify via History)",
        "scope": "whatsApp",
        "problem": "Was: default country code stored but not used. Verify saved code affects WhatsApp sends.",
    },
    "EC-AUTO-037": {
        "screen": "Billing cycle tab (number fields)",
        "scope": "billingCycle",
        "problem": "Clearing a number field may leave invalid empty value instead of resetting to min.",
    },
    "EC-AUTO-038": {
        "screen": "Payment reminders tab (max reminders field)",
        "scope": "paymentReminders",
        "problem": "Entering value above max may be accepted on client; save may fail without clear message.",
    },
    "EC-AUTO-043": {
        "screen": "WhatsApp tab (business hours start/end)",
        "scope": "whatsApp",
        "problem": "Invalid time format (e.g. 9:00) may pass client blur; server rejects on save.",
    },
    "EC-AUTO-044": {
        "screen": "WhatsApp tab (default country code)",
        "scope": "whatsApp",
        "problem": "Letters in country code may pass client; server rejects on save with toast only.",
    },
    "EC-AUTO-049": {
        "screen": "Automation Center — page error banner",
        "scope": "overview",
        "problem": "Network failure on config load may not offer working Retry.",
    },
    "EC-AUTO-050": {
        "screen": "Any config tab (save toast)",
        "scope": "whatsApp",
        "problem": "Validation errors from API may show generic 'Save failed' instead of field-specific message.",
    },
    "EC-AUTO-053": {
        "screen": "Overview (Run Now toast) + Run history",
        "scope": "overview",
        "problem": "Was: partial run showed success toast. Verify warning toast when status is PARTIAL.",
    },
    "EC-AUTO-054": {
        "screen": "Any config tab after Save",
        "scope": "invoices",
        "problem": "UI may not update after save without full page reload.",
    },
    "EC-AUTO-055": {
        "screen": "Overview → Run history tab",
        "scope": "history",
        "problem": "New run may not appear in History immediately after Run Now.",
    },
    "EC-AUTO-062": {
        "screen": "Master Settings (platform TZ) + Automation Overview (timezone field)",
        "scope": "overview",
        "problem": "Was: automation timezone not synced from platform settings. Verify sync or drift banner.",
    },
    "EC-AUTO-070": {
        "screen": "Overview (Run Now toast / dry run banner)",
        "scope": "overview",
        "problem": "Was: dry-run mode forced real Run Now silently. Verify dryRunForced warning in toast.",
    },
    "EC-AUTO-086": {
        "screen": "Overview run + Billing follow-ups (WhatsApp queue)",
        "scope": "overview",
        "problem": "Without WhatsApp API creds, reminders queue follow-up only — verify expected behaviour.",
    },
    "EC-AUTO-113": {
        "screen": "Overview or any config tab (Save / Run buttons)",
        "scope": "overview",
        "problem": "Double-click Save or Run may fire duplicate API requests.",
    },
    "EC-AUTO-118": {
        "screen": "Overview (Run Now)",
        "scope": "overview",
        "problem": "Mass email via Run Now has no confirmation dialog — risk/audit UX review needed.",
    },
    "EC-AUTO-121": {
        "screen": "Run history tab (error list)",
        "scope": "history",
        "problem": "Business names in run errors visible to admins — privacy/PII review needed.",
    },
    "EC-AUTO-126": {
        "screen": "Invoices tab (with config error banner)",
        "scope": "invoices",
        "problem": "Config load error banner may disappear when switching away from Overview tab.",
    },
    "EC-AUTO-128": {
        "screen": "Automation Center — all tabs (375px mobile)",
        "scope": "overview",
        "problem": "Layout may break or be unusable on small mobile screens.",
    },
    "EC-AUTO-129": {
        "screen": "Any config tab (Save) or Overview (Run Now)",
        "scope": "overview",
        "problem": "OPEN: Offline save/run shows generic error — needs clear 'You are offline' message.",
    },
    "EC-AUTO-130": {
        "screen": "Automation Center — scope sidebar + browser history",
        "scope": "invoices",
        "problem": "Was: browser back/forward lost selected tab. Verify ?scope= preserved.",
    },
    "EC-AUTO-131": {
        "screen": "Run history tab (long error text)",
        "scope": "history",
        "problem": "Was: long error lists overflow layout. Verify text wraps readably.",
    },
    "EC-AUTO-132": {
        "screen": "Automation Center — left scope sidebar (tabs)",
        "scope": "overview",
        "problem": "OPEN: Incomplete ARIA tab semantics — keyboard/screen reader navigation broken.",
    },
}

HEADERS = [
    "Test Case ID",
    "Screen / Tab (Where to Test)",
    "Scope URL Parameter",
    "Problem Description (What Is Wrong or Unverified)",
    "Category",
    "Feature Area",
    "Priority",
    "Severity",
    "Action Required",
    "Matrix Status",
    "Login Role Required",
    "Environment / Setup",
    "Full Page URL",
    "How to Test (Step by Step)",
    "Expected Outcome",
    "Pass Criteria (What You Should See)",
    "Watch For During Test",
    "Actual Result (Fill In)",
    "Pass / Fail (Fill In)",
    "Tester Name (Fill In)",
    "Test Date (Fill In)",
    "Notes (Fill In)",
]

ACTION_FILLS = {
    "Execute now": PatternFill("solid", fgColor="FFEB9C"),
    "Fix first, then test": PatternFill("solid", fgColor="FFC7CE"),
    "Regression verify": PatternFill("solid", fgColor="C6EFCE"),
}

PRIORITY_FILLS = {
    "P1": PatternFill("solid", fgColor="FFC7CE"),
    "P2": PatternFill("solid", fgColor="FFEB9C"),
    "P3": PatternFill("solid", fgColor="E2EFDA"),
}


def load_matrix_rows():
    matrix_path = Path(__file__).resolve().parent / "generate-automation-qa-excel.py"
    spec = importlib.util.spec_from_file_location("automation_matrix", matrix_path)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module.ROWS


def layer_from_module(module: str) -> str:
    if "(Frontend)" in module:
        return "Frontend"
    if "(Full-stack)" in module:
        return "Full-stack"
    return "Other"


def feature_area(module: str) -> str:
    name = module.split("(")[0].strip()
    return name


def requires_manual_frontend_test(row: dict) -> bool:
    if row["status"] == "N/A":
        return False
    layer = layer_from_module(row["module"])
    if layer == "Frontend":
        return True
    if layer != "Full-stack":
        return False
    tool = row["tool"]
    return row["type"] in ("Manual", "Both") or "Playwright" in tool


def login_role(pre: str, scenario: str) -> str:
    text = f"{pre} {scenario}".upper()
    if "MASTER_ADMIN" in text or "MASTER ADMIN" in text:
        return "Master Admin"
    if "PLATFORM_ADMIN" in text and "WITHOUT BILLING" in text:
        return "Platform Admin (no billing permission)"
    if "PLATFORM_ADMIN" in text or "PLATFORM ADMIN" in text:
        return "Platform Admin (with billing)"
    if "STORE OWNER" in text or "STAFF" in text:
        return "Store Owner or Staff"
    if "NO ACTIVE" in text or "WITHOUT LOGIN" in text or "INCOGNITO" in text:
        return "Not logged in (incognito)"
    if "EXPIRED" in text or "INVALID SESSION" in text:
        return "Expired / invalid session"
    return "Master Admin (default)"


def action_required(row: dict) -> str:
    if row["status"] == "Open":
        return "Fix first, then test"
    if row["status"] == "Not Tested":
        return "Execute now"
    if row["status"] == "Closed":
        return "Regression verify"
    return "Execute now"


def infer_screen_and_scope(row: dict) -> tuple[str, str]:
    case_id = row["id"]
    if case_id in CASE_SCREEN_AND_PROBLEM:
        entry = CASE_SCREEN_AND_PROBLEM[case_id]
        return entry["screen"], entry.get("scope", "")

    text = f"{row['scenario']} {row['steps']} {row['pre']}".lower()

    scope_rules: list[tuple[str, str, str]] = [
        ("history", "Run history tab", "history"),
        ("invoices", "Invoices tab", "invoices"),
        ("payment reminder", "Payment reminders tab", "paymentReminders"),
        ("follow-up", "Follow-ups tab", "followUps"),
        ("follow up", "Follow-ups tab", "followUps"),
        ("expiry", "Expiry & renewal tab", "expiryRenewal"),
        ("renewal", "Expiry & renewal tab", "expiryRenewal"),
        ("monthly report", "Monthly reports tab", "monthlyReports"),
        ("whatsapp", "WhatsApp tab", "whatsApp"),
        ("business hours", "WhatsApp tab", "whatsApp"),
        ("country code", "WhatsApp tab", "whatsApp"),
        ("billing cycle", "Billing cycle tab", "billingCycle"),
        ("grace period", "Billing cycle tab", "billingCycle"),
        ("run now", "Overview tab", "overview"),
        ("dry run", "Overview tab", "overview"),
        ("master switch", "Overview tab", "overview"),
        ("global.enabled", "Overview tab", "overview"),
        ("timezone", "Overview tab", "overview"),
        ("loading", "Automation Center — initial load", "overview"),
        ("config api", "Automation Center — page error banner", "overview"),
        ("network", "Automation Center — page error banner", "overview"),
        ("mobile", "Automation Center — all tabs (mobile)", "overview"),
        ("scope", "Automation Center — scope sidebar", "overview"),
        ("tab", "Automation Center — scope sidebar", "overview"),
        ("read-only", "Automation Center — all config tabs", "overview"),
        ("login", "Login page", ""),
        ("redirect", "Login / access guard", ""),
        ("incognito", "Login page", ""),
        ("session", "Login page / Automation Center", ""),
        ("master settings", "Master Settings + Overview timezone", "overview"),
        ("platform timezone", "Master Settings + Overview timezone", "overview"),
        ("toast", "Overview (toast notification)", "overview"),
    ]

    for keyword, screen, scope in scope_rules:
        if keyword in text:
            return screen, scope

    if "Authentication" in row["module"]:
        return "Login / access guard (before or on Automation page)", ""

    return "Automation Center — Overview tab", "overview"


def problem_description(row: dict) -> str:
    case_id = row["id"]
    if case_id in CASE_SCREEN_AND_PROBLEM:
        return CASE_SCREEN_AND_PROBLEM[case_id]["problem"]

    if row["status"] == "Open":
        return row["remarks"] or row["actual"] or f"Open defect: {row['scenario']}"

    if row["remarks"]:
        return row["remarks"]

    if row["status"] == "Not Tested":
        return (
            f"Not yet manually verified. Confirm on screen: {row['scenario']}. "
            f"Expected: {row['expected']}"
        )

    if row["status"] == "Closed":
        return (
            f"Previously fixed — regression check. Confirm: {row['expected']}"
        )

    return row["actual"] or row["scenario"]


def scope_url(scope: str) -> str:
    if not scope:
        return "(none — use base automation URL or login page)"
    label = SCOPE_LABELS.get(scope, scope)
    return f"?scope={scope}  ({label} tab)"


def full_page_url(scope: str) -> str:
    if not scope:
        return AUTOMATION_URL
    return f"{AUTOMATION_URL}?scope={scope}"


def known_issue(row: dict) -> str:
    if row["status"] == "Not Tested":
        return "Not yet executed — follow steps and record result"
    return "Fixed in code — confirm behaviour still matches expected outcome"


def watch_for(row: dict) -> str:
    if row["status"] == "Open":
        return "This case is OPEN — fail until problem column issue is resolved."
    if row["auto_notes"]:
        return row["auto_notes"]
    return known_issue(row)


def expand_steps(row: dict, role: str, scope: str) -> str:
    steps = row["steps"].strip()
    page_url = full_page_url(scope) if scope else AUTOMATION_URL
    lines = [
        f"1. Open browser (Chrome recommended). Base URL: {BASE_URL}",
        f"2. Log in as: {role} (or stay logged out if test requires it).",
        f"3. Navigate to: {page_url}",
    ]
    if scope:
        lines.append(
            f"4. Confirm left sidebar shows '{SCOPE_LABELS.get(scope, scope)}' tab selected."
        )
        step_offset = 5
    else:
        step_offset = 4

    step_num = step_offset
    for part in re.split(r"(?<=\d)\.\s+", steps):
        part = part.strip()
        if not part:
            continue
        part = re.sub(r"^\d+\.\s*", "", part)
        if part:
            lines.append(f"{step_num}. {part}")
            step_num += 1

    tool = row["tool"]
    if "offline" in row["scenario"].lower() or "network" in row["scenario"].lower():
        lines.append(
            f"{step_num}. Use DevTools → Network → Offline (or block API in Network tab) to simulate failure."
        )
        step_num += 1
    if "mobile" in row["scenario"].lower() or "375" in row["pre"]:
        lines.append(
            f"{step_num}. Set DevTools device toolbar to 375px width (iPhone SE) or real phone."
        )
        step_num += 1
    if "Playwright" in tool and row["type"] != "Manual":
        lines.append(
            f"{step_num}. Optional: same flow can be automated later with Playwright e2e."
        )

    return "\n".join(lines)


def pass_criteria(row: dict) -> str:
    expected = row["expected"]
    extras = []
    scenario_lower = row["scenario"].lower()

    if "toast" in scenario_lower:
        extras.append("Toast notification appears with correct title and readable message.")
    if "redirect" in expected.lower():
        extras.append("URL changes away from automation page or to login as specified.")
    if "disabled" in expected.lower():
        extras.append("Control is visibly disabled (greyed out) and cannot be clicked.")
    if "read-only" in expected.lower() or "disabled" in row["expected"].lower():
        extras.append("Save / Run buttons hidden or disabled; read-only hint visible if applicable.")
    if "history" in scenario_lower:
        extras.append("History tab lists runs with status badge, timestamp, and summary line.")
    if "refresh" in scenario_lower or "back/forward" in scenario_lower:
        extras.append("Same scope/tab remains selected after browser refresh or back/forward.")
    if "error" in scenario_lower or "fail" in scenario_lower:
        extras.append("Error is visible to user (banner or toast), not only in browser console.")
    if "loading" in scenario_lower:
        extras.append("Loading text or spinner visible before content appears.")

    if extras:
        return expected + "\n\nAlso verify:\n• " + "\n• ".join(extras)
    return expected


def environment_setup(row: dict) -> str:
    pre = row["pre"].strip()
    extras = []
    pre_lower = pre.lower()

    if "smtp" in pre_lower:
        extras.append("SMTP configured in .env.local OR intentionally unset to test skip/error paths.")
    if "global.enabled = false" in pre_lower or "global.enabled=false" in pre_lower:
        extras.append("In Automation → Overview: turn OFF 'Enable automations' master switch.")
    if "dryrun" in pre_lower.replace(" ", "") or "dry run" in pre_lower:
        extras.append("In Automation → Overview: turn ON 'Dry run mode' if test requires it.")
    if ">20" in pre or "20 automation runs" in pre_lower:
        extras.append("Seed or run automation 21+ times so History has more than 20 entries.")
    if "platform admin" in pre_lower and "billing" in pre_lower:
        extras.append("Use a Platform Admin account; billing permission on/off as stated in preconditions.")
    if "no prior automation runs" in pre_lower:
        extras.append("Use fresh test DB or delete AutomationRunLog rows for empty history test.")

    base = pre if pre else "Standard local dev environment running on port 3002."
    if extras:
        return base + "\n\nAdditional setup:\n• " + "\n• ".join(extras)
    return base


def row_to_values(row: dict) -> list:
    role = login_role(row["pre"], row["scenario"])
    screen, scope = infer_screen_and_scope(row)
    return [
        row["id"],
        screen,
        scope_url(scope),
        problem_description(row),
        row["module"],
        feature_area(row["module"]),
        row["priority"],
        row["severity"],
        action_required(row),
        row["status"],
        role,
        environment_setup(row),
        full_page_url(scope) if scope else AUTOMATION_URL,
        expand_steps(row, role, scope),
        row["expected"],
        pass_criteria(row),
        watch_for(row),
        "",
        "",
        "",
        "",
        "",
    ]


def write_data_sheet(ws, rows: list[dict], title_note: str = "") -> None:
    if title_note:
        ws.cell(row=1, column=1, value=title_note).font = Font(bold=True, size=12)
        start_row = 3
    else:
        start_row = 1

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF", size=11)

    for col, header in enumerate(HEADERS, 1):
        cell = ws.cell(row=start_row, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for idx, row in enumerate(rows, start_row + 1):
        for col_idx, value in enumerate(row_to_values(row), 1):
            cell = ws.cell(row=idx, column=col_idx, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if col_idx == 7:
                cell.fill = PRIORITY_FILLS.get(row["priority"], PatternFill())
            if col_idx == 9:
                action = action_required(row)
                cell.fill = ACTION_FILLS.get(action, PatternFill())
            if col_idx == 4 and row["status"] == "Open":
                cell.fill = PatternFill("solid", fgColor="FFC7CE")

    ws.freeze_panes = ws.cell(row=start_row + 1, column=1).coordinate
    if rows:
        last_row = start_row + len(rows)
        ws.auto_filter.ref = (
            f"A{start_row}:{get_column_letter(len(HEADERS))}{last_row}"
        )

    widths = [
        14, 36, 28, 44, 32, 22, 8, 10, 18, 12, 28, 36, 44, 48, 36, 40, 32, 20, 12, 16, 14, 24,
    ]
    for col_idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def write_instructions(ws) -> None:
    lines = [
        ("Automation Center — Manual QA Test Pack", 14, True),
        ("", 11, False),
        ("Purpose", 12, True),
        (
            "Execute every test case in the 'Execute now' and 'Fix first' sheets. "
            "Columns 'Screen / Tab' and 'Problem Description' tell you WHERE and WHAT. "
            "Fill: Actual Result, Pass/Fail, Tester Name, Test Date, Notes.",
            11,
            False,
        ),
        ("", 11, False),
        ("Environment", 12, True),
        (f"• Local app: {BASE_URL}", 11, False),
        (f"• Automation page: {AUTOMATION_URL}", 11, False),
        ("• Use Master Admin unless 'Login Role' column says otherwise", 11, False),
        ("• Do NOT run destructive tests against production/staging Supabase", 11, False),
        ("", 11, False),
        ("Automation Center screens (left sidebar tabs)", 12, True),
        ("• Overview — master switch, dry run, timezone, Run Now", 11, False),
        ("• Billing cycle — cycle start, due day, grace period", 11, False),
        ("• Invoices — auto-send, payment confirmation, sendOnRenewalDue", 11, False),
        ("• Payment reminders — email/WhatsApp reminder days", 11, False),
        ("• Follow-ups — cadence, max follow-ups, escalate", 11, False),
        ("• Expiry & renewal — renewal/expiry warning days", 11, False),
        ("• Monthly reports — schedule, recipients, summaries", 11, False),
        ("• WhatsApp — country code, business hours", 11, False),
        ("• Run history — past runs, badges, errors, Load more", 11, False),
        ("", 11, False),
        ("Key columns", 12, True),
        ("• Screen / Tab — exact UI area to open", 11, False),
        ("• Problem Description — what is wrong or what you must verify", 11, False),
        ("• Scope URL Parameter — add ?scope=... to jump to the tab", 11, False),
        ("", 11, False),
        ("Pass / Fail rules", 12, True),
        ("• PASS = Expected outcome and pass criteria fully met", 11, False),
        ("• FAIL = Any wrong redirect, missing UI, wrong toast, or broken control", 11, False),
        ("• BLOCKED = Cannot test (missing account, env, or open code fix)", 11, False),
        ("", 11, False),
        ("Sheet guide", 12, True),
        ("• Execute Now — cases you must run first (Not Tested)", 11, False),
        ("• Fix First — 2 cases need code fix before pass (129, 132)", 11, False),
        ("• All Manual Tests — complete list for full regression", 11, False),
        ("• By Category — grouped by feature area", 11, False),
    ]
    for row_idx, (text, size, bold) in enumerate(lines, 1):
        cell = ws.cell(row=row_idx, column=1, value=text)
        cell.font = Font(bold=bold, size=size)
    ws.column_dimensions["A"].width = 90


def write_summary(ws, all_rows: list[dict], execute: list[dict], fix_first: list[dict]) -> None:
    ws.cell(row=1, column=1, value="Summary").font = Font(bold=True, size=14)
    rows_data = [
        ("Total manual frontend test cases", len(all_rows)),
        ("Execute now (Not Tested)", len(execute)),
        ("Fix first, then test (Open)", len(fix_first)),
        ("Regression verify (Closed)", len([r for r in all_rows if r["status"] == "Closed"])),
    ]
    for i, (label, val) in enumerate(rows_data, 3):
        ws.cell(row=i, column=1, value=label).font = Font(bold=True)
        ws.cell(row=i, column=2, value=val)
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 12


def main() -> None:
    all_rows = load_matrix_rows()
    manual_rows = [r for r in all_rows if requires_manual_frontend_test(r)]
    manual_rows.sort(key=lambda r: r["id"])

    execute_now = [r for r in manual_rows if r["status"] == "Not Tested"]
    fix_first = [r for r in manual_rows if r["status"] == "Open"]
    regression = [r for r in manual_rows if r["status"] == "Closed"]

    wb = Workbook()
    instructions = wb.active
    instructions.title = "How To Use"
    write_instructions(instructions)

    summary = wb.create_sheet("Summary")
    write_summary(summary, manual_rows, execute_now, fix_first)

    execute_sheet = wb.create_sheet("Execute Now")
    write_data_sheet(
        execute_sheet,
        execute_now,
        f"RUN THESE FIRST — {len(execute_now)} tests (Status: Not Tested)",
    )

    fix_sheet = wb.create_sheet("Fix First")
    write_data_sheet(
        fix_sheet,
        fix_first,
        f"FIX BEFORE PASS — {len(fix_first)} tests (Status: Open)",
    )

    all_sheet = wb.create_sheet("All Manual Tests")
    write_data_sheet(
        all_sheet,
        manual_rows,
        f"COMPLETE MANUAL LIST — {len(manual_rows)} tests",
    )

    reg_sheet = wb.create_sheet("Regression Verify")
    write_data_sheet(
        reg_sheet,
        regression,
        f"OPTIONAL REGRESSION — {len(regression)} tests (Status: Closed)",
    )

    categories: dict[str, list[dict]] = {}
    for row in manual_rows:
        cat = feature_area(row["module"])
        categories.setdefault(cat, []).append(row)

    for cat in sorted(categories):
        safe_name = cat[:31].replace("/", "-")
        sheet = wb.create_sheet(safe_name)
        write_data_sheet(sheet, categories[cat])

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)
    print(f"Wrote {OUTPUT}")
    print(f"  All manual frontend tests: {len(manual_rows)}")
    print(f"  Execute now: {len(execute_now)}")
    print(f"  Fix first: {len(fix_first)}")
    print(f"  Regression verify: {len(regression)}")


if __name__ == "__main__":
    main()
