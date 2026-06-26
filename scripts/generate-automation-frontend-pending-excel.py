"""Export frontend EC-AUTO cases still pending execution (Not Tested + Open)."""
from __future__ import annotations

import importlib.util
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUTPUT = (
    Path(__file__).resolve().parent.parent
    / "docs"
    / "EC-AUTO-Frontend-Pending.xlsx"
)

HEADERS = [
    "Edge Case ID",
    "Layer",
    "Status",
    "Priority",
    "Severity",
    "Scenario Description",
    "Preconditions",
    "Test Steps",
    "Expected Result",
    "Known Issue / Gap",
    "Code Review",
    "Test Type",
    "Automation Tool",
    "Automation Notes",
    "Remarks",
]

STATUS_FILLS = {
    "Not Tested": PatternFill("solid", fgColor="FFEB9C"),
    "Open": PatternFill("solid", fgColor="FFC7CE"),
}

TYPE_FILLS = {
    "Auto": PatternFill("solid", fgColor="C6EFCE"),
    "Manual": PatternFill("solid", fgColor="FFEB9C"),
    "Both": PatternFill("solid", fgColor="BDD7EE"),
}


def load_matrix_rows():
    matrix_path = Path(__file__).resolve().parent / "generate-automation-qa-excel.py"
    spec = importlib.util.spec_from_file_location("automation_matrix", matrix_path)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module.ROWS


def layer_from_module(module: str) -> str:
    if "(Frontend)" in module:
        return "Frontend"
    if "(Full-stack)" in module:
        return "Full-stack"
    if "(Backend)" in module:
        return "Backend"
    return "Other"


def issue_text(row: dict) -> str:
    if row["status"] == "Open":
        return row["remarks"] or row["actual"] or "Not implemented — fix required before test"
    if row["remarks"]:
        return row["remarks"]
    if row["review"] == "Pass" and row["status"] == "Not Tested":
        return "Code review pass — pending execution / QA sign-off"
    return row["actual"] or "Pending test execution"


def is_frontend_pending(row: dict) -> bool:
    layer = layer_from_module(row["module"])
    return layer in ("Frontend", "Full-stack") and row["status"] in ("Not Tested", "Open")


def row_values(row: dict) -> list:
    return [
        row["id"],
        layer_from_module(row["module"]),
        row["status"],
        row["priority"],
        row["severity"],
        row["scenario"],
        row["pre"],
        row["steps"],
        row["expected"],
        issue_text(row),
        row["review"],
        row["type"],
        row["tool"],
        row["auto_notes"],
        row["remarks"],
    ]


def write_sheet(ws, rows: list[dict]) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF", size=11)

    for col, header in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_idx, row in enumerate(rows, 2):
        for col_idx, value in enumerate(row_values(row), 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if col_idx == 3:
                cell.fill = STATUS_FILLS.get(row["status"], PatternFill())
            if col_idx == 12:
                cell.fill = TYPE_FILLS.get(row["type"], PatternFill())

    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{len(rows) + 1}"

    widths = [14, 12, 12, 8, 10, 42, 34, 42, 34, 40, 10, 10, 22, 28, 28]
    for col_idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def write_summary(ws, rows: list[dict]) -> None:
    from collections import Counter

    by_layer = Counter(layer_from_module(r["module"]) for r in rows)
    by_status = Counter(r["status"] for r in rows)

    ws.cell(row=1, column=1, value="Frontend EC-AUTO — pending execution").font = Font(
        bold=True, size=14
    )
    ws.cell(row=3, column=1, value="Total pending").font = Font(bold=True)
    ws.cell(row=3, column=2, value=len(rows))

    row = 5
    ws.cell(row=row, column=1, value="By status").font = Font(bold=True)
    row += 1
    for label in ("Open", "Not Tested"):
        if by_status[label]:
            ws.cell(row=row, column=1, value=label)
            ws.cell(row=row, column=2, value=by_status[label])
            row += 1

    row += 1
    ws.cell(row=row, column=1, value="By layer").font = Font(bold=True)
    row += 1
    for label in ("Frontend", "Full-stack"):
        if by_layer[label]:
            ws.cell(row=row, column=1, value=label)
            ws.cell(row=row, column=2, value=by_layer[label])
            row += 1

    row += 2
    ws.cell(row=row, column=1, value="Open = fix required before test").font = Font(italic=True)
    row += 1
    ws.cell(row=row, column=1, value="Not Tested = execute Playwright / manual QA").font = Font(
        italic=True
    )

    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 12


def main() -> None:
    all_rows = load_matrix_rows()
    pending = [row for row in all_rows if is_frontend_pending(row)]
    pending.sort(key=lambda row: (row["status"] != "Open", row["id"]))

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    write_summary(summary, pending)

    all_sheet = wb.create_sheet("All Pending")
    write_sheet(all_sheet, pending)

    for layer in ("Frontend", "Full-stack"):
        layer_rows = [row for row in pending if layer_from_module(row["module"]) == layer]
        if layer_rows:
            sheet = wb.create_sheet(layer)
            write_sheet(sheet, layer_rows)

    open_rows = [row for row in pending if row["status"] == "Open"]
    if open_rows:
        sheet = wb.create_sheet("Open (fix first)")
        write_sheet(sheet, open_rows)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)
    print(f"Wrote {OUTPUT} ({len(pending)} frontend pending cases)")


if __name__ == "__main__":
    main()
